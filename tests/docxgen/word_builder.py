from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_word_safely(data) -> str:
    """
    Safely compiles a .docx file using python-docx based on structured schema data.
    Enforces directory locking to prevent path traversal attacks.
    """
    # 🛡️ Directory Lock Security Check
    script_dir = Path(__file__).parent.resolve()
    clean_filename = Path(data.filename).name
    if not clean_filename.endswith(".docx"):
        clean_filename += ".docx"

    target_path = (script_dir / clean_filename).resolve()
    if target_path.parent != script_dir:
        raise PermissionError("Security Violation: Attempted directory traversal outside folder.")

    # Create Document
    doc = docx.Document()

    # 📄 Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 🎨 Title Paragraph
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run(data.document_title)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 78, 121) # Dark Blue

    # Subtitle / Author
    if data.author_or_subtitle:
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(18)
        sub_run = sub_p.add_run(data.author_or_subtitle)
        sub_run.font.name = "Calibri"
        sub_run.font.size = Pt(12)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 116, 139) # Slate Grey

    # Add Content Sections
    for sec in data.sections:
        # Heading 2
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(12)
        h_p.paragraph_format.space_after = Pt(4)
        h_run = h_p.add_run(sec.heading)
        h_run.font.name = "Calibri"
        h_run.font.size = Pt(14)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(15, 23, 42)

        # Body Paragraphs
        for para_text in sec.paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            p_run = p.add_run(para_text)
            p_run.font.name = "Calibri"
            p_run.font.size = Pt(11)
            p_run.font.color.rgb = RGBColor(51, 65, 85)

        # Bullet Points (if any)
        if sec.bullet_points:
            for bullet in sec.bullet_points:
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_after = Pt(3)
                b_run = bp.add_run(bullet)
                b_run.font.name = "Calibri"
                b_run.font.size = Pt(11)
                b_run.font.color.rgb = RGBColor(51, 65, 85)

    # Save document
    doc.save(target_path)
    return str(target_path)