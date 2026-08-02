from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_safely(data) -> str:
    """
    Safely compiles an .xlsx file using openpyxl based on structured schema data.
    Enforces directory locking to prevent path traversal attacks.
    """
    # 🛡️ Directory Lock Security Check
    script_dir = Path(__file__).parent.resolve()
    clean_filename = Path(data.filename).name
    if not clean_filename.endswith(".xlsx"):
        clean_filename += ".xlsx"

    target_path = (script_dir / clean_filename).resolve()
    if target_path.parent != script_dir:
        raise PermissionError("Security Violation: Attempted directory traversal outside folder.")

    # Create Workbook
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 🎨 Styling Configs
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Dark Blue
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    data_font = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for sheet_data in data.sheets:
        ws = wb.create_sheet(title=sheet_data.sheet_name[:31]) # Max 31 chars for Excel sheet names

        # Add Title Row
        ws.cell(row=1, column=1, value=data.title).font = title_font
        ws.row_dimensions[1].height = 25

        # Add Headers
        start_row = 3
        for col_num, header_title in enumerate(sheet_data.headers, 1):
            cell = ws.cell(row=start_row, column=col_num, value=header_title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        ws.row_dimensions[start_row].height = 20

        # Add Data Rows
        for row_idx, row_values in enumerate(sheet_data.rows, start=start_row + 1):
            for col_idx, val in enumerate(row_values, start=1):
                # Try parsing numeric values so Excel handles them as numbers rather than strings
                try:
                    if "." in str(val):
                        val_to_write = float(val)
                    else:
                        val_to_write = int(val)
                except ValueError:
                    val_to_write = val

                cell = ws.cell(row=row_idx, column=col_idx, value=val_to_write)
                cell.font = data_font
                cell.border = thin_border
                
                # Align numbers right, text left
                if isinstance(val_to_write, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(horizontal="left")

        # 📐 Auto-fit Column Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save workbook
    wb.save(target_path)
    return str(target_path)